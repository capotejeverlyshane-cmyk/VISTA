"""Verify citation numbering in the fixed document."""
import re
from docx import Document

doc = Document(r"VISTA_Chapter1_and_2_FIXED.docx")

all_citations = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if not text:
        continue
    matches = re.finditer(r'\[(\d+)\]', text)
    for m in matches:
        cite_num = int(m.group(1))
        start = max(0, m.start() - 50)
        end = min(len(text), m.end() + 50)
        snippet = text[start:end]
        all_citations.append((i, cite_num, snippet))

print("ALL CITATIONS IN FIXED DOCUMENT:")
print("-" * 100)
for para_idx, cite_num, snippet in all_citations:
    print(f"  Para {para_idx:3d} | [{cite_num:2d}] | ...{snippet}...")

used = sorted(set(c[1] for c in all_citations))
print(f"\nUnique numbers used: {used}")
expected = list(range(1, max(used) + 1))
missing = [n for n in expected if n not in used]
if missing:
    print(f"GAPS FOUND: {missing}")
else:
    print("NO GAPS - citation numbering is clean!")

# Print references section
print("\nREFERENCES:")
print("-" * 100)
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if re.match(r'^\[\d+\]', text):
        print(f"  {text[:140]}")
