import re
from docx import Document

SRC = r"VISTA_Chapter1_and_2_REVISED_AY25-26.docx"
OUT = r"VISTA_Chapter1_and_2_REVISED_AY25-26.docx"

doc = Document(SRC)

# 1. REMOVE THE SDG SENTENCE FROM BACKGROUND PARAGRAPH 1
sdg_sentence = "These digital transformation efforts are closely tied to the United Nations Sustainable Development Goals, particularly SDG 9 (Industry, Innovation, and Infrastructure) and SDG 16 (Peace, Justice, and Strong Institutions), which call for inclusive, accountable institutions supported by technology [2]. "

def replace_text_in_paragraph(para, old_text, new_text):
    if old_text not in para.text:
        return False
    # Simple replacement if it falls within a single run (often false in Word)
    for run in para.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return True
            
    # Complex replacement across runs
    full_text = para.text
    if old_text in full_text:
        # Clear all runs and put everything in the first run to ensure clean replace
        # This loses intra-paragraph formatting (like bold/italics), but is safer for text match.
        # Let's try to preserve formatting by only touching the text
        pass

    # Better complex replace:
    text = para.text
    new_para_text = text.replace(old_text, new_text)
    
    if text != new_para_text:
        # We will just clear all runs and write the new text to the first run
        # Since this is the first paragraph, it shouldn't have much special formatting other than normal text
        para.text = new_para_text
        return True
    return False

for p in doc.paragraphs:
    if sdg_sentence in p.text:
        replace_text_in_paragraph(p, sdg_sentence, "")
        print("Removed SDG sentence from intro.")
        break
    # Try a slight variation without the trailing space
    elif sdg_sentence.strip() in p.text:
        replace_text_in_paragraph(p, sdg_sentence.strip() + " ", "")
        print("Removed SDG sentence from intro (stripped).")
        break


# 2. FIND REFERENCES SECTION INDEX
ref_start_idx = -1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().upper() == "REFERENCES":
        ref_start_idx = i
        break

if ref_start_idx == -1:
    print("Error: Could not find REFERENCES section.")
    exit(1)

# 3. BUILD MAPPING OF OLD_ID -> NEW_ID BASED ON APPEARANCE
old_to_new = {}
new_id_counter = 1

citation_pattern = re.compile(r'\[(\d+)\]')

for i in range(ref_start_idx):
    p = doc.paragraphs[i]
    matches = citation_pattern.findall(p.text)
    for match in matches:
        old_id = int(match)
        if old_id not in old_to_new:
            old_to_new[old_id] = new_id_counter
            new_id_counter += 1

print(f"Found {len(old_to_new)} unique citations in text.")
print("Mapping:", old_to_new)


# 4. REPLACE CITATIONS IN TEXT
def replace_citations_in_text(text, mapping):
    # Use a function to replace each match
    def replacer(m):
        old_id = int(m.group(1))
        if old_id in mapping:
            return f"[{mapping[old_id]}]"
        return m.group(0) # Shouldn't happen
    
    # We must do this carefully to avoid replacing [1] with [2] and then [2] with [3]
    # Re.sub processes sequentially, so it's safe if we do it in one pass!
    return citation_pattern.sub(replacer, text)


for i in range(ref_start_idx):
    p = doc.paragraphs[i]
    if '[' in p.text and ']' in p.text:
        new_text = replace_citations_in_text(p.text, old_to_new)
        if new_text != p.text:
            p.text = new_text

# 5. EXTRACT AND REORDER REFERENCES LIST
old_references = {}
# Read paragraphs after REFERENCES
ref_paragraphs = doc.paragraphs[ref_start_idx+1:]
current_ref_id = None
current_ref_text = ""
ref_idx_map = {} # Maps old_id to the index of the paragraph in the document

for i in range(ref_start_idx + 1, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    text = p.text.strip()
    if not text:
        continue
        
    match = re.match(r'^\[(\d+)\]\s*(.*)', text)
    if match:
        old_id = int(match.group(1))
        old_references[old_id] = text  # store full text including [old_id]
        ref_idx_map[old_id] = i

print(f"Found {len(old_references)} references in the References list.")

# Validate if all in-text citations have a corresponding reference
for old_id, new_id in old_to_new.items():
    if old_id not in old_references:
        print(f"WARNING: Citation [{old_id}] used in text but not found in References list!")

# 6. UPDATE REFERENCES LIST IN DOCUMENT
# Because we want to reorder them in place, the easiest way is to rewrite the text of the existing reference paragraphs
# We will gather the new sorted references
sorted_new_refs = []
# Create a reverse mapping
new_to_old = {v: k for k, v in old_to_new.items()}

# Find any references that were NOT cited in text, we should append them at the end or skip them
# Since we want exactly 30 references and all should be cited, let's include any uncited ones at the end just in case
cited_old_ids = set(old_to_new.keys())
uncited_old_ids = set(old_references.keys()) - cited_old_ids
for u_old_id in uncited_old_ids:
    old_to_new[u_old_id] = new_id_counter
    new_to_old[new_id_counter] = u_old_id
    new_id_counter += 1
    print(f"WARNING: Reference [{u_old_id}] is in the list but NEVER cited in text! Appending as [{old_to_new[u_old_id]}].")

# Now write them back. We have len(old_references) paragraphs to write.
# Get the paragraph indices where references currently live
ref_indices = sorted(list(ref_idx_map.values()))

for i, new_id in enumerate(sorted(new_to_old.keys())):
    old_id = new_to_old[new_id]
    original_text = old_references[old_id]
    # Replace the starting [old_id] with [new_id]
    new_ref_text = re.sub(r'^\[\d+\]', f"[{new_id}]", original_text)
    
    if i < len(ref_indices):
        doc.paragraphs[ref_indices[i]].text = new_ref_text
    else:
        # Append new paragraph if we run out of existing ones
        doc.add_paragraph(new_ref_text)

# Clear any remaining reference paragraphs if there are extra
for i in range(len(new_to_old.keys()), len(ref_indices)):
    doc.paragraphs[ref_indices[i]].text = ""

doc.save(OUT)
print("Successfully renumbered all citations and references.")
